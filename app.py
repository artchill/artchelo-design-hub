from flask import (
    Flask, render_template, request, redirect, url_for, flash, abort, send_file
)
from flask_login import (
    LoginManager, login_user, logout_user, login_required, current_user
)
from werkzeug.security import check_password_hash, generate_password_hash
from datetime import datetime, date, timedelta
from calendar import monthrange, month_name
from functools import wraps
from pathlib import Path
from urllib.parse import urlparse
import csv
import io
import re
import time
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from sqlalchemy import extract, func
from models import (
    db, User, Attendance, LeaveDay,
    Client, SalesLead, SalesInvoice,
    Project, DomainRenewal, ProjectStatusHistory,
)

app = Flask(__name__)
app.config['SECRET_KEY'] = 'change-this-secret-key-in-production'
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///attendance.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'


# ---------- Helpers ----------

TIER_RANK = {
    'Bronze': 1, 'Essential': 1,
    'Silver': 2, 'Advanced': 2,
    'Gold': 3, 'Elite': 3,
}


GENDERS = ('male', 'female', 'prefer_not_to_say')
EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')

AVATAR_DIR = Path(app.root_path) / 'static' / 'uploads' / 'avatars'
AVATAR_MAX_BYTES = 2 * 1024 * 1024
AVATAR_MIME_TYPES = ('image/jpeg', 'image/png', 'image/webp')
AVATAR_EXTENSIONS = {'JPEG': 'jpg', 'PNG': 'png', 'WEBP': 'webp'}


def validate_profile_fields(codename, email, gender, exclude_user_id=None):
    """Validates codename/email/gender. Returns a list of error messages."""
    errors = []
    if codename and len(codename) > 30:
        errors.append('Codename must be 30 characters or fewer.')
    if email and not EMAIL_RE.match(email):
        errors.append('Invalid email address.')
    elif email:
        query = User.query.filter(User.email == email)
        if exclude_user_id is not None:
            query = query.filter(User.id != exclude_user_id)
        if query.first():
            errors.append(f'Email "{email}" is already in use.')
    if gender and gender not in GENDERS:
        errors.append('Invalid gender selected.')
    return errors


def read_avatar_upload(file):
    """Returns (image_bytes, error). image_bytes is None when nothing was uploaded."""
    if not file or not file.filename:
        return None, None
    if file.mimetype not in AVATAR_MIME_TYPES:
        return None, 'Profile picture must be a JPG, PNG, or WEBP image.'
    data = file.read(AVATAR_MAX_BYTES + 1)
    if len(data) > AVATAR_MAX_BYTES:
        return None, 'Profile picture must be 2MB or smaller.'
    try:
        with Image.open(io.BytesIO(data)) as img:
            if img.format not in AVATAR_EXTENSIONS:
                return None, 'Profile picture must be a JPG, PNG, or WEBP image.'
            img.verify()
    except Exception:
        return None, 'Profile picture is not a valid image.'
    return data, None


def delete_avatar_file(filename):
    if filename:
        (AVATAR_DIR / Path(filename).name).unlink(missing_ok=True)


def save_avatar(user, data):
    """Resizes to max 400x400, saves to disk, and replaces the user's old avatar file."""
    AVATAR_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.open(io.BytesIO(data))
    ext = AVATAR_EXTENSIONS[img.format]
    img = ImageOps.exif_transpose(img)
    img.thumbnail((400, 400))
    filename = f'avatar_{user.id}_{int(time.time())}.{ext}'
    img.save(AVATAR_DIR / filename)
    if user.profile_pic != filename:
        delete_avatar_file(user.profile_pic)
    user.profile_pic = filename


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


def has_role(user, role):
    """Super-admin passes 'admin' checks; every other role must match exactly."""
    if role == 'admin':
        return user.is_admin_or_above
    return user.role == role


def assignable_roles(user):
    """Roles `user` may give to others when creating or editing an account."""
    if user.is_super_admin:
        return ('employee', 'admin', 'super-admin')
    return ('employee', 'admin')


def safe_referrer(fallback):
    """request.referrer if it's on this host and not the current page, else fallback."""
    ref = request.referrer
    if ref and urlparse(ref).netloc == request.host and ref != request.url:
        return ref
    return fallback


def role_required(role):
    def decorator(f):
        @wraps(f)
        def wrapped(*args, **kwargs):
            if not current_user.is_authenticated:
                return redirect(url_for('login'))
            if not has_role(current_user, role):
                if role == 'super-admin' and current_user.is_admin_or_above:
                    flash('Access denied: this action requires super-admin privileges.', 'error')
                    return redirect(safe_referrer(url_for('dashboard')))
                flash('Access denied.', 'error')
                return redirect(url_for('login'))
            return f(*args, **kwargs)
        return wrapped
    return decorator


def generate_invoice_number():
    year = datetime.now().year
    prefix = f'INV-{year}-'
    last = (SalesInvoice.query
            .filter(SalesInvoice.invoice_number.like(f'{prefix}%'))
            .order_by(SalesInvoice.id.desc())
            .first())
    if last:
        try:
            last_num = int(last.invoice_number.split('-')[-1])
            return f'{prefix}{last_num + 1:03d}'
        except (ValueError, IndexError):
            pass
    return f'{prefix}001'


def build_month_calendar(user_id, year, month):
    days_in_month = monthrange(year, month)[1]
    today = date.today()
    start = date(year, month, 1)
    end = date(year, month, days_in_month)

    attendances = Attendance.query.filter(
        Attendance.user_id == user_id,
        Attendance.datetime_in >= datetime.combine(start, datetime.min.time()),
        Attendance.datetime_in <= datetime.combine(end, datetime.max.time()),
    ).all()
    att_by_date = {}
    for a in attendances:
        att_by_date.setdefault(a.datetime_in.date(), []).append(a)

    leaves = LeaveDay.query.filter(
        LeaveDay.user_id == user_id,
        LeaveDay.date >= start,
        LeaveDay.date <= end,
    ).all()
    leave_by_date = {l.date: l for l in leaves}

    calendar = []
    for day_num in range(1, days_in_month + 1):
        d = date(year, month, day_num)
        is_weekend = d.weekday() >= 5
        is_future = d > today
        day_attendances = att_by_date.get(d, [])
        leave = leave_by_date.get(d)

        if day_attendances:
            status = 'present'
        elif leave:
            status = leave.leave_type.lower()
        elif is_future:
            status = 'future'
        elif is_weekend:
            status = 'no_work'
        else:
            status = 'absent'

        calendar.append({
            'date': d, 'day_name': d.strftime('%A'), 'status': status,
            'attendances': day_attendances, 'leave': leave,
            'is_future': is_future, 'is_weekend': is_weekend,
        })

    return calendar


# ---------- Auth ----------

@app.route('/')
def index():
    if current_user.is_authenticated:
        return redirect(url_for(
            'dashboard' if current_user.is_admin_or_above
            else 'employee_dashboard'
        ))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('index'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password_hash, password):
            if user.employment_status and user.employment_status != 'active':
                status_label = user.employment_status.replace('_', ' ').title()
                flash(f'Account not active (status: {status_label}). Please contact admin.', 'error')
                return render_template('login.html')
            login_user(user)
            return redirect(url_for(
                'dashboard' if user.is_admin_or_above
                else 'employee_dashboard'
            ))
        flash('Invalid username or password.', 'error')
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))


@app.route('/account/change-password', methods=['GET', 'POST'])
@login_required
def change_password():
    if request.method == 'POST':
        current_password = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        confirm_password = request.form.get('confirm_password', '')

        errors = []
        if not current_password or not new_password or not confirm_password:
            errors.append('All fields are required.')
        if not check_password_hash(current_user.password_hash, current_password):
            errors.append('Current password is incorrect.')
        if len(new_password) < 6:
            errors.append('New password must be at least 6 characters.')
        if new_password == current_password:
            errors.append('New password must be different from current password.')
        if new_password != confirm_password:
            errors.append('New password and confirmation do not match.')

        if errors:
            for e in errors:
                flash(e, 'error')
            return render_template('change_password.html')

        current_user.password_hash = generate_password_hash(new_password, method='pbkdf2:sha256')
        db.session.commit()
        flash('Password changed successfully.', 'success')
        return redirect(url_for(
            'dashboard' if current_user.is_admin_or_above else 'employee_dashboard'
        ))

    return render_template('change_password.html')


# ---------- Employee ----------

@app.route('/employee/dashboard')
@role_required('employee')
def employee_dashboard():
    today = date.today()
    year = int(request.args.get('year', today.year))
    month = int(request.args.get('month', today.month))

    open_session = Attendance.query.filter_by(
        user_id=current_user.id, datetime_out=None
    ).first()
    calendar = build_month_calendar(current_user.id, year, month)

    return render_template(
        'employee_dashboard.html',
        open_session=open_session, calendar=calendar,
        year=year, month=month, month_name=month_name[month],
        available_years=range(today.year - 2, today.year + 2),
    )


@app.route('/employee/time-in', methods=['POST'])
@role_required('employee')
def time_in():
    existing = Attendance.query.filter_by(
        user_id=current_user.id, datetime_out=None
    ).first()
    if existing:
        flash('You already have an active session.', 'error')
        return redirect(url_for('employee_dashboard'))
    now = datetime.now()
    record = Attendance(
        user_id=current_user.id, datetime_in=now, day=now.strftime('%A'),
    )
    db.session.add(record)
    db.session.commit()
    flash(f'Timed in at {now.strftime("%I:%M %p")}.', 'success')
    return redirect(url_for('employee_dashboard'))


@app.route('/employee/time-out', methods=['POST'])
@role_required('employee')
def time_out():
    task_report = request.form.get('task_report', '').strip()
    if not task_report:
        flash('A task report is required to time out.', 'error')
        return redirect(url_for('employee_dashboard'))
    session_ = Attendance.query.filter_by(
        user_id=current_user.id, datetime_out=None
    ).first()
    if not session_:
        flash('No active session found.', 'error')
        return redirect(url_for('employee_dashboard'))
    now = datetime.now()
    session_.datetime_out = now
    session_.total_hours = round(
        (now - session_.datetime_in).total_seconds() / 3600, 2
    )
    session_.task_report = task_report
    db.session.commit()
    flash(f'Timed out. Total hours: {session_.total_hours}', 'success')
    return redirect(url_for('employee_dashboard'))


# ---------- Leave marking ----------

@app.route('/leave/mark', methods=['POST'])
@login_required
def mark_leave():
    target_user_id = int(request.form.get('user_id', current_user.id))
    date_str = request.form.get('date', '').strip()
    leave_type = request.form.get('leave_type', '').strip().upper()
    note = request.form.get('note', '').strip()
    redirect_to = request.form.get('redirect_to', 'employee_dashboard')
    redirect_kwargs = {}
    if redirect_to == 'admin_employee_tracker':
        redirect_kwargs['employee_id'] = target_user_id
        redirect_kwargs['year'] = request.form.get('year')
        redirect_kwargs['month'] = request.form.get('month')

    if not current_user.is_admin_or_above and target_user_id != current_user.id:
        flash('You can only mark leave for yourself.', 'error')
        return redirect(url_for(redirect_to, **redirect_kwargs))
    if leave_type not in ('VL', 'SL'):
        flash('Leave type must be VL or SL.', 'error')
        return redirect(url_for(redirect_to, **redirect_kwargs))
    try:
        leave_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        flash('Invalid date.', 'error')
        return redirect(url_for(redirect_to, **redirect_kwargs))

    has_att = Attendance.query.filter(
        Attendance.user_id == target_user_id,
        Attendance.datetime_in >= datetime.combine(leave_date, datetime.min.time()),
        Attendance.datetime_in <= datetime.combine(leave_date, datetime.max.time()),
    ).first()
    if has_att:
        flash('Cannot mark leave — attendance already recorded.', 'error')
        return redirect(url_for(redirect_to, **redirect_kwargs))

    existing = LeaveDay.query.filter_by(
        user_id=target_user_id, date=leave_date
    ).first()
    if existing:
        existing.leave_type = leave_type
        existing.note = note or None
        existing.marked_by_id = current_user.id
        existing.marked_at = datetime.now()
        flash(f'Updated {leave_date} to {leave_type}.', 'success')
    else:
        new_leave = LeaveDay(
            user_id=target_user_id, date=leave_date,
            leave_type=leave_type, note=note or None,
            marked_by_id=current_user.id,
        )
        db.session.add(new_leave)
        flash(f'Marked {leave_date} as {leave_type}.', 'success')
    db.session.commit()
    return redirect(url_for(redirect_to, **redirect_kwargs))


@app.route('/leave/<int:leave_id>/delete', methods=['POST'])
@login_required
def delete_leave(leave_id):
    leave = LeaveDay.query.get_or_404(leave_id)
    if not current_user.is_admin_or_above and leave.user_id != current_user.id:
        flash('Not allowed.', 'error')
        return redirect(url_for('employee_dashboard'))
    redirect_to = request.form.get('redirect_to', 'employee_dashboard')
    redirect_kwargs = {}
    if redirect_to == 'admin_employee_tracker':
        redirect_kwargs['employee_id'] = leave.user_id
        redirect_kwargs['year'] = request.form.get('year')
        redirect_kwargs['month'] = request.form.get('month')
    db.session.delete(leave)
    db.session.commit()
    flash('Leave entry removed.', 'success')
    return redirect(url_for(redirect_to, **redirect_kwargs))


# ---------- Admin: Attendance ----------

@app.route('/admin/dashboard')
@role_required('admin')
def admin_dashboard():
    first_employee = User.query.filter_by(role='employee').order_by(User.fname).first()
    if first_employee:
        return redirect(url_for('admin_employee_tracker',
                                employee_id=first_employee.id))
    return render_template('admin_dashboard_empty.html')


@app.route('/admin/employees/<int:employee_id>/tracker')
@role_required('admin')
def admin_employee_tracker(employee_id):
    employee = User.query.get_or_404(employee_id)
    today = date.today()
    year = int(request.args.get('year', today.year))
    month = int(request.args.get('month', today.month))
    calendar = build_month_calendar(employee_id, year, month)
    all_employees = User.query.filter_by(role='employee').order_by(User.fname).all()
    return render_template(
        'admin_employee_tracker.html',
        employee=employee, all_employees=all_employees, calendar=calendar,
        year=year, month=month, month_name=month_name[month],
        available_years=range(today.year - 2, today.year + 2),
    )


# ---------- Attendance Export ----------

EXPORT_STATUS_LABELS = {
    'present': 'Present', 'vl': 'VL', 'sl': 'SL',
    'absent': 'Absent', 'no_work': 'Weekend', 'future': 'Future',
}
EXPORT_COLUMNS = ['Date', 'Day', 'Status', 'Time In', 'Time Out', 'Hours', 'Task / Note']


def _export_year_month():
    try:
        year = int(request.args['year'])
        month = int(request.args['month'])
    except (KeyError, ValueError):
        abort(400, 'year and month query parameters are required integers.')
    if not 1 <= month <= 12:
        abort(400, 'month must be between 1 and 12.')
    return year, month


def build_attendance_export(employee, year, month):
    """Returns (header_rows, table_rows, summary_rows, sheet_title, filename_stem).
    Table rows hold raw values: hours is a float or None."""
    calendar = build_month_calendar(employee.id, year, month)

    header = [
        ('Employee', employee.full_name),
        ('Codename', f'@{employee.codename}' if employee.codename else '—'),
        ('Position', employee.job_position or '—'),
        ('Email', employee.email),
        ('Month', f'{month_name[month]} {year}'),
        ('Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
    ]

    rows = []
    counts = {label: 0 for label in EXPORT_STATUS_LABELS.values()}
    total_hours = 0.0
    for day in calendar:
        status = EXPORT_STATUS_LABELS.get(day['status'], day['status'].upper())
        counts[status] = counts.get(status, 0) + 1
        time_in = time_out = ''
        hours = None
        note = ''

        sessions = day['attendances']
        if sessions:
            time_in = min(a.datetime_in for a in sessions).strftime('%I:%M %p')
            outs = [a.datetime_out for a in sessions if a.datetime_out]
            if outs:
                time_out = max(outs).strftime('%I:%M %p')
            logged = [a.total_hours for a in sessions if a.total_hours is not None]
            if logged:
                hours = sum(logged)
                total_hours += hours
            note = ' | '.join(a.task_report for a in sessions if a.task_report)
        elif day['leave'] and day['leave'].note:
            note = day['leave'].note

        rows.append([
            day['date'].strftime('%Y-%m-%d'), day['day_name'], status,
            time_in, time_out, hours, note,
        ])

    summary = [
        ('Total Present days', counts['Present']),
        ('Total Absent days', counts['Absent']),
        ('Total VL days', counts['VL']),
        ('Total SL days', counts['SL']),
        ('Total Weekend days', counts['Weekend']),
        ('Total Hours Worked', round(total_hours, 2)),
    ]

    sheet_title = f'{month_name[month]} {year}'
    stem = f'attendance_{employee.fname}_{employee.lname}_{year}_{month:02d}'
    stem = stem.lower().replace(' ', '_')
    return header, rows, summary, sheet_title, stem


@app.route('/admin/employees/<int:employee_id>/export/xlsx')
@role_required('admin')
def export_attendance_xlsx(employee_id):
    employee = User.query.get_or_404(employee_id)
    year, month = _export_year_month()
    header, rows, summary, sheet_title, stem = build_attendance_export(employee, year, month)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    bold = Font(bold=True)

    for label, value in header:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])

    ws.append(EXPORT_COLUMNS)
    table_header_row = ws.max_row
    header_fill = PatternFill('solid', fgColor='F1F5F9')
    header_border = Border(bottom=Side(style='thin', color='94A3B8'))
    for cell in ws[table_header_row]:
        cell.font = bold
        cell.fill = header_fill
        cell.border = header_border

    status_fills = {
        'Present': 'DCFCE7', 'VL': 'F3E8FF', 'SL': 'FEF3C7',
        'Absent': 'FEE2E2', 'Weekend': 'F1F5F9',
    }
    for row in rows:
        ws.append(row)
        r = ws.max_row
        color = status_fills.get(row[2])
        if color:
            ws.cell(row=r, column=3).fill = PatternFill('solid', fgColor=color)
        ws.cell(row=r, column=6).number_format = '0.00'

    ws.append([])
    for label, value in summary:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=2).number_format = '0.00'

    # Column A also holds the header/summary labels, so it's wider than a date needs.
    widths = {'A': 20, 'B': 12, 'C': 10, 'D': 10, 'E': 10, 'F': 8, 'G': 50}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=table_header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name=f'{stem}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/admin/employees/<int:employee_id>/export/csv')
@role_required('admin')
def export_attendance_csv(employee_id):
    employee = User.query.get_or_404(employee_id)
    year, month = _export_year_month()
    header, rows, summary, _, stem = build_attendance_export(employee, year, month)

    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerows(header)
    writer.writerow([])
    writer.writerow(EXPORT_COLUMNS)
    for row in rows:
        hours = row[5]
        writer.writerow(row[:5] + ['' if hours is None else f'{hours:.2f}'] + row[6:])
    writer.writerow([])
    writer.writerows(summary[:-1])
    label, total = summary[-1]
    writer.writerow([label, f'{total:.2f}'])

    buf = io.BytesIO(out.getvalue().encode('utf-8-sig'))
    return send_file(
        buf, as_attachment=True, download_name=f'{stem}.csv',
        mimetype='text/csv',
    )


@app.route('/admin/employees')
@role_required('admin')
def admin_employees():
    show_all = request.args.get('show_all', '') == '1'
    search = request.args.get('q', '').strip()
    query = User.query
    if not show_all:
        query = query.filter(User.employment_status == 'active')
    if search:
        like = f'%{search}%'
        query = query.filter(db.or_(
            User.fname.ilike(like),
            User.lname.ilike(like),
        ))
    employees = query.order_by(User.fname).all()
    return render_template('admin_employees.html',
                           employees=employees, show_all=show_all, search=search)


@app.route('/admin/employees/new', methods=['GET', 'POST'])
@role_required('admin')
def admin_new_employee():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        confirm_password = request.form.get('confirm_password', '')
        fname = request.form.get('fname', '').strip()
        mname = request.form.get('mname', '').strip().rstrip('.').upper()[:1] or None
        lname = request.form.get('lname', '').strip()
        role = request.form.get('role', '').strip()
        job_position = request.form.get('job_position', '').strip()
        birthdate_str = request.form.get('birthdate', '').strip()
        date_hired_str = request.form.get('date_hired', '').strip()
        monthly_salary_str = request.form.get('monthly_salary', '').strip()
        employee_type = request.form.get('employee_type', '').strip()
        codename = request.form.get('codename', '').strip() or None
        email = request.form.get('email', '').strip().lower()
        gender = request.form.get('gender', '').strip()

        errors = []
        if not all([username, password, fname, lname, role, job_position,
                    birthdate_str, date_hired_str, monthly_salary_str, employee_type,
                    email, gender]):
            errors.append('All fields are required.')
        errors.extend(validate_profile_fields(codename, email, gender))
        avatar_data, avatar_error = read_avatar_upload(request.files.get('profile_pic'))
        if avatar_error:
            errors.append(avatar_error)
        if role == 'super-admin' and not current_user.is_super_admin:
            errors.append('Access denied: only super-admin can assign super-admin role.')
        elif role not in assignable_roles(current_user):
            errors.append('Invalid role selected.')
        if employee_type and employee_type not in ('full-time', 'part-time', 'outsourced'):
            errors.append('Invalid employee type.')
        if len(password) < 6:
            errors.append('Password must be at least 6 characters.')
        if password != confirm_password:
            errors.append('Passwords do not match.')
        if User.query.filter_by(username=username).first():
            errors.append(f'Username "{username}" is already taken.')

        birthdate = None
        date_hired = None
        monthly_salary = None
        try:
            birthdate = datetime.strptime(birthdate_str, '%Y-%m-%d').date()
        except (ValueError, TypeError):
            if birthdate_str:
                errors.append('Invalid birthdate.')
        try:
            date_hired = datetime.strptime(date_hired_str, '%Y-%m-%d').date()
        except (ValueError, TypeError):
            if date_hired_str:
                errors.append('Invalid date hired.')
        try:
            monthly_salary = float(monthly_salary_str) if monthly_salary_str else None
            if monthly_salary is not None and monthly_salary < 0:
                errors.append('Salary must be a positive number.')
        except ValueError:
            errors.append('Invalid salary.')

        if errors:
            for e in errors:
                flash(e, 'error')
            return render_template('admin_new_employee.html',
                form_data={'username': username, 'fname': fname, 'mname': mname or '',
                           'lname': lname, 'role': role, 'job_position': job_position,
                           'birthdate': birthdate_str, 'date_hired': date_hired_str,
                           'monthly_salary': monthly_salary_str,
                           'employee_type': employee_type,
                           'codename': codename or '', 'email': email, 'gender': gender})

        new_user = User(
            username=username,
            password_hash=generate_password_hash(password, method='pbkdf2:sha256'),
            fname=fname, mname=mname, lname=lname, role=role,
            job_position=job_position,
            birthdate=birthdate, date_hired=date_hired,
            monthly_salary=monthly_salary,
            employee_type=employee_type,
            employment_status='active',
            codename=codename, email=email, gender=gender,
        )
        db.session.add(new_user)
        db.session.commit()
        if avatar_data:
            # Needs new_user.id for the filename, so save after the first commit
            save_avatar(new_user, avatar_data)
            db.session.commit()
        flash(f'User "{username}" created successfully.', 'success')
        return redirect(url_for('admin_employees'))
    return render_template('admin_new_employee.html', form_data={})


@app.route('/admin/employees/<int:user_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def admin_edit_employee(user_id):
    emp = User.query.get_or_404(user_id)
    # An admin editing a super-admin can't touch role, password or status:
    # a password reset would let them log in as the super-admin.
    protected = emp.is_super_admin and not current_user.is_super_admin

    if request.method == 'POST':
        codename = request.form.get('codename', '').strip() or None
        email = request.form.get('email', '').strip().lower()
        gender = request.form.get('gender', '').strip()
        new_role = request.form.get('role', emp.role)

        errors = []
        if new_role != emp.role:
            if protected:
                errors.append('Access denied: only super-admin can change a super-admin\'s role.')
            elif new_role == 'super-admin' and not current_user.is_super_admin:
                errors.append('Access denied: only super-admin can assign super-admin role')
            elif new_role not in assignable_roles(current_user):
                errors.append('Invalid role selected.')
            elif emp.is_super_admin and User.query.filter_by(role='super-admin').count() <= 1:
                errors.append('Cannot change the role of the last super-admin.')
        if protected and (request.form.get('new_password', '').strip()
                          or request.form.get('employment_status', emp.employment_status) != emp.employment_status):
            errors.append('Access denied: only super-admin can reset a super-admin\'s password or status.')
        if not email or not gender:
            errors.append('Email and gender are required.')
        errors.extend(validate_profile_fields(codename, email, gender, exclude_user_id=emp.id))
        avatar_data, avatar_error = read_avatar_upload(request.files.get('profile_pic'))
        if avatar_error:
            errors.append(avatar_error)
        if errors:
            for e in errors:
                flash(e, 'error')
            return render_template('admin_edit_employee.html', emp=emp)

        emp.codename = codename
        emp.email = email
        emp.gender = gender

        emp.fname = request.form.get('fname', '').strip() or emp.fname
        emp.mname = (request.form.get('mname', '').strip().rstrip('.').upper()[:1] or None)
        emp.lname = request.form.get('lname', '').strip() or emp.lname
        emp.username = request.form.get('username', '').strip() or emp.username
        emp.role = new_role
        emp.job_position = request.form.get('job_position', '').strip() or emp.job_position
        emp.employee_type = request.form.get('employee_type', emp.employee_type)

        birthdate_str = request.form.get('birthdate', '').strip()
        if birthdate_str:
            try:
                emp.birthdate = datetime.strptime(birthdate_str, '%Y-%m-%d').date()
            except ValueError:
                pass

        date_hired_str = request.form.get('date_hired', '').strip()
        if date_hired_str:
            try:
                emp.date_hired = datetime.strptime(date_hired_str, '%Y-%m-%d').date()
            except ValueError:
                pass

        salary_str = request.form.get('monthly_salary', '').strip()
        if salary_str:
            try:
                emp.monthly_salary = float(salary_str)
            except ValueError:
                pass

        new_status = request.form.get('employment_status', emp.employment_status)
        if new_status != emp.employment_status:
            if new_status in ('resigned', 'terminated'):
                resign_str = request.form.get('resignation_date', '').strip()
                if not resign_str:
                    flash('Resignation date is required for Resigned or Terminated status.', 'error')
                    return render_template('admin_edit_employee.html', emp=emp)
                try:
                    emp.resignation_date = datetime.strptime(resign_str, '%Y-%m-%d').date()
                except ValueError:
                    flash('Invalid resignation date.', 'error')
                    return render_template('admin_edit_employee.html', emp=emp)
            elif new_status == 'active':
                emp.resignation_date = None
            emp.employment_status = new_status

        new_password = request.form.get('new_password', '').strip()
        if new_password:
            if len(new_password) < 6:
                flash('Password must be at least 6 characters.', 'error')
                return render_template('admin_edit_employee.html', emp=emp)
            emp.password_hash = generate_password_hash(new_password, method='pbkdf2:sha256')
            flash('Password updated.', 'success')

        # Touch the disk last so an early return above can't orphan or delete files
        if avatar_data:
            save_avatar(emp, avatar_data)
        elif request.form.get('remove_avatar'):
            delete_avatar_file(emp.profile_pic)
            emp.profile_pic = None

        db.session.commit()
        flash(f'Profile for {emp.full_name} updated.', 'success')
        return redirect(url_for('admin_employees'))

    return render_template('admin_edit_employee.html', emp=emp)


@app.route('/admin/attendance/<int:record_id>/edit', methods=['GET', 'POST'])
@role_required('super-admin')
def admin_edit_attendance(record_id):
    record = Attendance.query.get_or_404(record_id)
    if request.method == 'POST':
        date_str = request.form.get('date', '').strip()
        time_in_str = request.form.get('time_in', '').strip()
        time_out_str = request.form.get('time_out', '').strip()
        task_report = request.form.get('task_report', '').strip()
        errors = []
        if not date_str or not time_in_str:
            errors.append('Date and Time In are required.')
        new_datetime_in = None
        new_datetime_out = None
        if not errors:
            try:
                new_datetime_in = datetime.strptime(
                    f'{date_str} {time_in_str}', '%Y-%m-%d %H:%M')
            except ValueError:
                errors.append('Invalid Date or Time In format.')
        if time_out_str and not errors:
            try:
                new_datetime_out = datetime.strptime(
                    f'{date_str} {time_out_str}', '%Y-%m-%d %H:%M')
                if new_datetime_out <= new_datetime_in:
                    errors.append('Time Out must be after Time In.')
            except ValueError:
                errors.append('Invalid Time Out format.')
        if errors:
            for e in errors:
                flash(e, 'error')
            return render_template('admin_edit_attendance.html', record=record)
        record.datetime_in = new_datetime_in
        record.day = new_datetime_in.strftime('%A')
        if new_datetime_out:
            record.datetime_out = new_datetime_out
            record.total_hours = round(
                (new_datetime_out - new_datetime_in).total_seconds() / 3600, 2)
        else:
            record.datetime_out = None
            record.total_hours = None
        record.task_report = task_report or None
        record.edited_at = datetime.now()
        record.edited_by_id = current_user.id
        db.session.commit()
        flash('Attendance record updated successfully.', 'success')
        return redirect(url_for('admin_employee_tracker',
                                employee_id=record.user_id))
    return render_template('admin_edit_attendance.html', record=record)


# ========== SALES MODULE ==========

@app.route('/sales')
@role_required('admin')
def sales_overview():
    today = date.today()
    month_start = date(today.year, today.month, 1)

    pipeline_value = db.session.query(func.sum(SalesLead.estimated_value)).filter(
        SalesLead.status.in_(('new', 'contacted', 'quoted'))
    ).scalar() or 0

    revenue_this_month = db.session.query(func.sum(SalesInvoice.amount)).filter(
        SalesInvoice.status == 'paid',
        SalesInvoice.paid_date >= month_start,
    ).scalar() or 0

    outstanding = db.session.query(func.sum(SalesInvoice.amount)).filter(
        SalesInvoice.status.in_(('sent', 'overdue'))
    ).scalar() or 0

    leads_won_month = SalesLead.query.filter(
        SalesLead.status == 'won',
        SalesLead.closed_at >= datetime.combine(month_start, datetime.min.time())
    ).count()

    recent_leads = SalesLead.query.order_by(SalesLead.created_at.desc()).limit(5).all()
    recent_invoices = SalesInvoice.query.order_by(SalesInvoice.issue_date.desc()).limit(5).all()

    return render_template(
        'sales_overview.html',
        pipeline_value=pipeline_value,
        revenue_this_month=revenue_this_month,
        outstanding=outstanding,
        leads_won_month=leads_won_month,
        recent_leads=recent_leads,
        recent_invoices=recent_invoices,
    )


# ----- Leads CRUD -----

@app.route('/sales/leads')
@role_required('admin')
def sales_leads():
    status_filter = request.args.get('status', '')
    query = SalesLead.query
    if status_filter:
        query = query.filter_by(status=status_filter)
    leads = query.order_by(SalesLead.created_at.desc()).all()
    return render_template('sales_leads.html', leads=leads, status_filter=status_filter)


@app.route('/sales/leads/new', methods=['GET', 'POST'])
@role_required('admin')
def sales_lead_new():
    if request.method == 'POST':
        lead = SalesLead(
            prospect_name=request.form.get('prospect_name', '').strip(),
            prospect_email=request.form.get('prospect_email', '').strip() or None,
            prospect_phone=request.form.get('prospect_phone', '').strip() or None,
            source=request.form.get('source', '').strip() or None,
            service_interest=request.form.get('service_interest', '').strip() or None,
            estimated_value=float(request.form.get('estimated_value') or 0) or None,
            status=request.form.get('status', 'new'),
            notes=request.form.get('notes', '').strip() or None,
        )
        followup = request.form.get('next_followup_date', '').strip()
        if followup:
            try:
                lead.next_followup_date = datetime.strptime(followup, '%Y-%m-%d').date()
            except ValueError:
                pass
        if not lead.prospect_name:
            flash('Prospect name is required.', 'error')
            return render_template('sales_lead_form.html', lead=lead, is_new=True)
        db.session.add(lead)
        db.session.commit()
        flash('Lead created.', 'success')
        return redirect(url_for('sales_leads'))
    return render_template('sales_lead_form.html', lead=None, is_new=True)


@app.route('/sales/leads/<int:lead_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def sales_lead_edit(lead_id):
    lead = SalesLead.query.get_or_404(lead_id)
    if request.method == 'POST':
        lead.prospect_name = request.form.get('prospect_name', '').strip()
        lead.prospect_email = request.form.get('prospect_email', '').strip() or None
        lead.prospect_phone = request.form.get('prospect_phone', '').strip() or None
        lead.source = request.form.get('source', '').strip() or None
        lead.service_interest = request.form.get('service_interest', '').strip() or None
        lead.estimated_value = float(request.form.get('estimated_value') or 0) or None
        new_status = request.form.get('status', 'new')
        if new_status in ('won', 'lost') and lead.status not in ('won', 'lost'):
            lead.closed_at = datetime.now()
        lead.status = new_status
        lead.notes = request.form.get('notes', '').strip() or None
        followup = request.form.get('next_followup_date', '').strip()
        if followup:
            try:
                lead.next_followup_date = datetime.strptime(followup, '%Y-%m-%d').date()
            except ValueError:
                lead.next_followup_date = None
        else:
            lead.next_followup_date = None
        db.session.commit()
        flash('Lead updated.', 'success')
        return redirect(url_for('sales_leads'))
    return render_template('sales_lead_form.html', lead=lead, is_new=False)


@app.route('/sales/leads/<int:lead_id>/delete', methods=['POST'])
@role_required('admin')
def sales_lead_delete(lead_id):
    lead = SalesLead.query.get_or_404(lead_id)
    db.session.delete(lead)
    db.session.commit()
    flash('Lead deleted.', 'success')
    return redirect(url_for('sales_leads'))


# ----- Invoices CRUD -----

@app.route('/sales/invoices')
@role_required('admin')
def sales_invoices():
    status_filter = request.args.get('status', '')
    query = SalesInvoice.query
    if status_filter:
        query = query.filter_by(status=status_filter)
    invoices = query.order_by(SalesInvoice.issue_date.desc()).all()
    return render_template('sales_invoices.html', invoices=invoices, status_filter=status_filter)


@app.route('/sales/invoices/new', methods=['GET', 'POST'])
@role_required('admin')
def sales_invoice_new():
    clients = Client.query.order_by(Client.name).all()
    if request.method == 'POST':
        client_id = int(request.form.get('client_id', 0) or 0)
        amount = float(request.form.get('amount') or 0)
        issue_date_str = request.form.get('issue_date', '').strip()
        if not client_id or amount <= 0 or not issue_date_str:
            flash('Client, amount, and issue date are required.', 'error')
            return render_template('sales_invoice_form.html',
                                   invoice=None, clients=clients, is_new=True)
        try:
            issue_dt = datetime.strptime(issue_date_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Invalid issue date.', 'error')
            return render_template('sales_invoice_form.html',
                                   invoice=None, clients=clients, is_new=True)

        due_date_str = request.form.get('due_date', '').strip()
        due_dt = None
        if due_date_str:
            try:
                due_dt = datetime.strptime(due_date_str, '%Y-%m-%d').date()
            except ValueError:
                pass

        invoice = SalesInvoice(
            invoice_number=generate_invoice_number(),
            client_id=client_id,
            amount=amount,
            issue_date=issue_dt,
            due_date=due_dt,
            payment_method=request.form.get('payment_method', '').strip() or None,
            status=request.form.get('status', 'draft'),
            description=request.form.get('description', '').strip() or None,
            notes=request.form.get('notes', '').strip() or None,
        )
        db.session.add(invoice)
        db.session.commit()
        flash(f'Invoice {invoice.invoice_number} created.', 'success')
        return redirect(url_for('sales_invoices'))
    return render_template('sales_invoice_form.html',
                           invoice=None, clients=clients, is_new=True)


@app.route('/sales/invoices/<int:invoice_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def sales_invoice_edit(invoice_id):
    invoice = SalesInvoice.query.get_or_404(invoice_id)
    clients = Client.query.order_by(Client.name).all()
    if request.method == 'POST':
        invoice.client_id = int(request.form.get('client_id', invoice.client_id))
        invoice.amount = float(request.form.get('amount') or invoice.amount)
        issue_date_str = request.form.get('issue_date', '').strip()
        if issue_date_str:
            try:
                invoice.issue_date = datetime.strptime(issue_date_str, '%Y-%m-%d').date()
            except ValueError:
                pass
        due_date_str = request.form.get('due_date', '').strip()
        if due_date_str:
            try:
                invoice.due_date = datetime.strptime(due_date_str, '%Y-%m-%d').date()
            except ValueError:
                invoice.due_date = None
        else:
            invoice.due_date = None

        new_status = request.form.get('status', invoice.status)
        if new_status == 'paid' and invoice.status != 'paid':
            paid_date_str = request.form.get('paid_date', '').strip()
            if paid_date_str:
                try:
                    invoice.paid_date = datetime.strptime(paid_date_str, '%Y-%m-%d').date()
                except ValueError:
                    invoice.paid_date = date.today()
            else:
                invoice.paid_date = date.today()
        elif new_status != 'paid':
            invoice.paid_date = None
        invoice.status = new_status

        invoice.payment_method = request.form.get('payment_method', '').strip() or None
        invoice.description = request.form.get('description', '').strip() or None
        invoice.notes = request.form.get('notes', '').strip() or None

        db.session.commit()
        flash(f'Invoice {invoice.invoice_number} updated.', 'success')
        return redirect(url_for('sales_invoices'))
    return render_template('sales_invoice_form.html',
                           invoice=invoice, clients=clients, is_new=False)


@app.route('/sales/invoices/<int:invoice_id>/delete', methods=['POST'])
@role_required('admin')
def sales_invoice_delete(invoice_id):
    invoice = SalesInvoice.query.get_or_404(invoice_id)
    db.session.delete(invoice)
    db.session.commit()
    flash('Invoice deleted.', 'success')
    return redirect(url_for('sales_invoices'))


# ----- Clients CRUD -----

@app.route('/sales/clients')
@role_required('admin')
def sales_clients():
    clients = Client.query.order_by(Client.name).all()
    return render_template('sales_clients.html', clients=clients)


@app.route('/sales/clients/new', methods=['GET', 'POST'])
@role_required('admin')
def sales_client_new():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if not name:
            flash('Client name is required.', 'error')
            return render_template('sales_client_form.html', client=None, is_new=True)
        client = Client(
            name=name,
            client_type=request.form.get('client_type', 'author'),
            email=request.form.get('email', '').strip() or None,
            phone=request.form.get('phone', '').strip() or None,
            notes=request.form.get('notes', '').strip() or None,
        )
        db.session.add(client)
        db.session.commit()
        flash(f'Client "{client.name}" added.', 'success')
        return redirect(url_for('sales_clients'))
    return render_template('sales_client_form.html', client=None, is_new=True)


@app.route('/sales/clients/<int:client_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def sales_client_edit(client_id):
    client = Client.query.get_or_404(client_id)
    if request.method == 'POST':
        client.name = request.form.get('name', '').strip() or client.name
        client.client_type = request.form.get('client_type', client.client_type)
        client.email = request.form.get('email', '').strip() or None
        client.phone = request.form.get('phone', '').strip() or None
        client.notes = request.form.get('notes', '').strip() or None
        db.session.commit()
        flash('Client updated.', 'success')
        return redirect(url_for('sales_clients'))
    return render_template('sales_client_form.html', client=client, is_new=False)


@app.route('/sales/clients/<int:client_id>/delete', methods=['POST'])
@role_required('admin')
def sales_client_delete(client_id):
    client = Client.query.get_or_404(client_id)
    if client.invoices:
        flash(f'Cannot delete {client.name} — client has invoices.', 'error')
        return redirect(url_for('sales_clients'))
    db.session.delete(client)
    db.session.commit()
    flash('Client deleted.', 'success')
    return redirect(url_for('sales_clients'))


# ========== MAIN DASHBOARD ==========

@app.route('/dashboard')
@role_required('admin')
def dashboard():
    today = date.today()
    month_start = date(today.year, today.month, 1)
    last_month_start = (month_start - timedelta(days=1)).replace(day=1)

    revenue_this_month = db.session.query(func.sum(SalesInvoice.amount)).filter(
        SalesInvoice.status == 'paid',
        SalesInvoice.paid_date >= month_start,
    ).scalar() or 0

    revenue_last_month = db.session.query(func.sum(SalesInvoice.amount)).filter(
        SalesInvoice.status == 'paid',
        SalesInvoice.paid_date >= last_month_start,
        SalesInvoice.paid_date < month_start,
    ).scalar() or 0

    pipeline_value = db.session.query(func.sum(SalesLead.estimated_value)).filter(
        SalesLead.status.in_(('new', 'contacted', 'quoted'))
    ).scalar() or 0

    outstanding = db.session.query(func.sum(SalesInvoice.amount)).filter(
        SalesInvoice.status.in_(('sent', 'overdue'))
    ).scalar() or 0

    active_clients = Client.query.count()

    leads_won_month = SalesLead.query.filter(
        SalesLead.status == 'won',
        SalesLead.closed_at >= datetime.combine(month_start, datetime.min.time())
    ).count()

    open_leads = SalesLead.query.filter(
        SalesLead.status.in_(('new', 'contacted', 'quoted'))
    ).count()

    projects_total = Project.query.count()
    projects_ongoing = Project.query.filter_by(status='ongoing').count()
    projects_completed = Project.query.filter_by(status='completed').count()
    projects_delivered_month = Project.query.filter(
        Project.delivered_date >= month_start
    ).count()

    recent_delivered = Project.query.filter(
        Project.delivered_date >= today - timedelta(days=90),
        Project.tat_days != None,
    ).all()
    avg_tat = (sum(p.tat_days for p in recent_delivered) / len(recent_delivered)
               if recent_delivered else 0)

    revenue_trend_labels = []
    revenue_trend_data = []
    for i in range(5, -1, -1):
        m = (month_start - timedelta(days=i * 30)).replace(day=1)
        next_m = (m + timedelta(days=32)).replace(day=1)
        rev = db.session.query(func.sum(SalesInvoice.amount)).filter(
            SalesInvoice.status == 'paid',
            SalesInvoice.paid_date >= m,
            SalesInvoice.paid_date < next_m,
        ).scalar() or 0
        revenue_trend_labels.append(m.strftime('%b %Y'))
        revenue_trend_data.append(float(rev))

    funnel_data = {
        'new': SalesLead.query.filter_by(status='new').count(),
        'contacted': SalesLead.query.filter_by(status='contacted').count(),
        'quoted': SalesLead.query.filter_by(status='quoted').count(),
        'won': SalesLead.query.filter_by(status='won').count(),
        'lost': SalesLead.query.filter_by(status='lost').count(),
    }

    invoice_status = {
        'draft': SalesInvoice.query.filter_by(status='draft').count(),
        'sent': SalesInvoice.query.filter_by(status='sent').count(),
        'paid': SalesInvoice.query.filter_by(status='paid').count(),
        'overdue': SalesInvoice.query.filter_by(status='overdue').count(),
        'void': SalesInvoice.query.filter_by(status='void').count(),
    }

    sources_query = db.session.query(
        SalesLead.source, func.count(SalesLead.id)
    ).group_by(SalesLead.source).all()
    sources_data = {s or 'Unknown': c for s, c in sources_query}

    top_clients_query = db.session.query(
        Client.name, func.sum(SalesInvoice.amount)
    ).join(SalesInvoice, SalesInvoice.client_id == Client.id).filter(
        SalesInvoice.status == 'paid'
    ).group_by(Client.name).order_by(func.sum(SalesInvoice.amount).desc()).limit(5).all()
    top_clients_labels = [name for name, _ in top_clients_query]
    top_clients_data = [float(amt) for _, amt in top_clients_query]

    tier_query = db.session.query(
        Project.package_tier, func.count(Project.id)
    ).group_by(Project.package_tier).all()
    tier_data = {t or 'Unknown': c for t, c in tier_query}

    workload_query = db.session.query(
        User.fname, func.count(Project.id)
    ).join(Project, Project.assigned_developer_id == User.id).filter(
        Project.status == 'ongoing'
    ).group_by(User.fname).all()
    workload_labels = [name for name, _ in workload_query]
    workload_data = [count for _, count in workload_query]

    recent_leads = SalesLead.query.order_by(SalesLead.created_at.desc()).limit(5).all()
    recent_invoices = SalesInvoice.query.order_by(SalesInvoice.issue_date.desc()).limit(5).all()
    recent_projects = Project.query.order_by(Project.created_at.desc()).limit(5).all()

    upcoming_followups = SalesLead.query.filter(
        SalesLead.next_followup_date != None,
        SalesLead.next_followup_date >= today,
        SalesLead.next_followup_date <= today + timedelta(days=7),
        SalesLead.status.in_(('new', 'contacted', 'quoted')),
    ).order_by(SalesLead.next_followup_date).all()

    return render_template(
        'dashboard.html',
        revenue_this_month=revenue_this_month,
        revenue_last_month=revenue_last_month,
        pipeline_value=pipeline_value,
        outstanding=outstanding,
        active_clients=active_clients,
        leads_won_month=leads_won_month,
        open_leads=open_leads,
        projects_total=projects_total,
        projects_ongoing=projects_ongoing,
        projects_completed=projects_completed,
        projects_delivered_month=projects_delivered_month,
        avg_tat=avg_tat,
        revenue_trend_labels=revenue_trend_labels,
        revenue_trend_data=revenue_trend_data,
        funnel_data=funnel_data,
        invoice_status=invoice_status,
        sources_data=sources_data,
        top_clients_labels=top_clients_labels,
        top_clients_data=top_clients_data,
        tier_data=tier_data,
        workload_labels=workload_labels,
        workload_data=workload_data,
        recent_leads=recent_leads,
        recent_invoices=recent_invoices,
        recent_projects=recent_projects,
        upcoming_followups=upcoming_followups,
        current_month_name=today.strftime('%B %Y'),
    )


# ========== WEBSITE PROJECTS MODULE ==========

@app.route('/projects')
@role_required('admin')
def projects_list():
    status_filter = request.args.get('status', '')
    tier_filter = request.args.get('tier', '')
    dev_filter = request.args.get('dev', '')
    search = request.args.get('q', '').strip()

    query = Project.query.join(Client, Project.client_id == Client.id)

    if status_filter:
        query = query.filter(Project.status == status_filter)
    if tier_filter:
        query = query.filter(Project.package_tier == tier_filter)
    if dev_filter:
        query = query.filter(Project.assigned_developer_id == int(dev_filter))
    if search:
        like = f'%{search}%'
        query = query.filter(db.or_(
            Client.name.ilike(like),
            Project.domain.ilike(like),
        ))

    projects = query.order_by(Project.created_at.desc()).all()
    all_devs = User.query.order_by(User.fname).all()

    return render_template(
        'projects_list.html',
        projects=projects, all_devs=all_devs,
        status_filter=status_filter, tier_filter=tier_filter,
        dev_filter=dev_filter, search=search,
    )


@app.route('/projects/new', methods=['GET', 'POST'])
@role_required('admin')
def project_new():
    clients = Client.query.order_by(Client.name).all()
    devs = User.query.order_by(User.fname).all()
    if request.method == 'POST':
        client_id = int(request.form.get('client_id', 0) or 0)
        if not client_id:
            flash('Client is required.', 'error')
            return render_template('project_form.html',
                                   project=None, clients=clients, devs=devs, is_new=True)

        tier = request.form.get('package_tier', '').strip() or None
        project = Project(
            client_id=client_id,
            domain=request.form.get('domain', '').strip().lower() or None,
            service_type=request.form.get('service_type', 'website'),
            package_tier=tier,
            tier_rank=TIER_RANK.get(tier),
            status=request.form.get('status', 'lead'),
            drive_url=request.form.get('drive_url', '').strip() or None,
            form_url=request.form.get('form_url', '').strip() or None,
            notes=request.form.get('notes', '').strip() or None,
        )
        dev_id = request.form.get('assigned_developer_id', '').strip()
        if dev_id:
            project.assigned_developer_id = int(dev_id)

        start_str = request.form.get('start_date', '').strip()
        if start_str:
            try:
                project.start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
                project.created_month = project.start_date.month
                project.created_year = project.start_date.year
            except ValueError:
                pass

        delivered_str = request.form.get('delivered_date', '').strip()
        if delivered_str:
            try:
                project.delivered_date = datetime.strptime(delivered_str, '%Y-%m-%d').date()
                if project.start_date:
                    project.tat_days = (project.delivered_date - project.start_date).days
            except ValueError:
                pass

        db.session.add(project)
        db.session.commit()

        db.session.add(ProjectStatusHistory(
            project_id=project.id, from_status=None,
            to_status=project.status, changed_by_id=current_user.id,
            note='Project created',
        ))
        db.session.commit()

        flash('Project created.', 'success')
        return redirect(url_for('projects_list'))
    return render_template('project_form.html',
                           project=None, clients=clients, devs=devs, is_new=True)


@app.route('/projects/<int:project_id>/edit', methods=['GET', 'POST'])
@role_required('admin')
def project_edit(project_id):
    project = Project.query.get_or_404(project_id)
    clients = Client.query.order_by(Client.name).all()
    devs = User.query.order_by(User.fname).all()

    if request.method == 'POST':
        old_status = project.status

        project.client_id = int(request.form.get('client_id', project.client_id))
        project.domain = request.form.get('domain', '').strip().lower() or None
        project.service_type = request.form.get('service_type', project.service_type)

        tier = request.form.get('package_tier', '').strip() or None
        project.package_tier = tier
        project.tier_rank = TIER_RANK.get(tier)

        project.status = request.form.get('status', project.status)

        dev_id = request.form.get('assigned_developer_id', '').strip()
        project.assigned_developer_id = int(dev_id) if dev_id else None

        project.drive_url = request.form.get('drive_url', '').strip() or None
        project.form_url = request.form.get('form_url', '').strip() or None
        project.notes = request.form.get('notes', '').strip() or None

        start_str = request.form.get('start_date', '').strip()
        if start_str:
            try:
                project.start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
                project.created_month = project.start_date.month
                project.created_year = project.start_date.year
            except ValueError:
                pass
        else:
            project.start_date = None

        delivered_str = request.form.get('delivered_date', '').strip()
        if delivered_str:
            try:
                project.delivered_date = datetime.strptime(delivered_str, '%Y-%m-%d').date()
                if project.start_date:
                    project.tat_days = (project.delivered_date - project.start_date).days
            except ValueError:
                pass
        else:
            project.delivered_date = None
            project.tat_days = None

        if old_status != project.status:
            db.session.add(ProjectStatusHistory(
                project_id=project.id, from_status=old_status,
                to_status=project.status, changed_by_id=current_user.id,
            ))

        db.session.commit()
        flash('Project updated.', 'success')
        return redirect(url_for('projects_list'))

    return render_template('project_form.html',
                           project=project, clients=clients, devs=devs, is_new=False)


@app.route('/projects/<int:project_id>/delete', methods=['POST'])
@role_required('admin')
def project_delete(project_id):
    project = Project.query.get_or_404(project_id)
    db.session.delete(project)
    db.session.commit()
    flash('Project deleted.', 'success')
    return redirect(url_for('projects_list'))


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True, port=5001)